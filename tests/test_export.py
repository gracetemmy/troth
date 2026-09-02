"""The workbook is built from memory at the moment it is requested."""

from datetime import date

from openpyxl import load_workbook

from troth import commitments as C
from troth import export as X
from troth import memory as M

TODAY = date(2026, 9, 3)


def _fixture(mem):
    M.remember_commitment(mem, "nimbus", "quote by Friday",
                          date(2026, 9, 4), "friday", C.EXACT, "Rep")
    M.remember_commitment(mem, "orbital", "paperwork by 15 September",
                          date(2026, 9, 15), "15 september", C.EXACT, "Rep")
    M.remember_commitment(mem, "orbital", "analytics pricing at some point",
                          None, "at some point", C.VAGUE, "Rep")
    return mem


def test_rows_sort_overdue_and_soonest_first(mem):
    _fixture(mem)
    rows = X.commitment_rows(mem, today=TODAY)
    assert [r["due"] for r in rows] == [date(2026, 9, 4), date(2026, 9, 15), None]


def test_date_filter_narrows(mem):
    _fixture(mem)
    assert len(X.commitment_rows(mem, None, date(2026, 9, 7), TODAY)) == 2


def test_undated_survives_every_filter(mem):
    """A commitment with no deadline is the one most likely to slip.
    Hiding it because it fails a date range would defeat the point."""
    _fixture(mem)
    for lo, hi in [(None, date(2026, 9, 4)), (date(2026, 9, 20), None)]:
        rows = X.commitment_rows(mem, lo, hi, TODAY)
        assert any(r["due"] is None for r in rows)


def test_workbook_scopes(mem, tmp_path):
    _fixture(mem)
    M.remember_client(mem, "nimbus", {"facts": ["signed in March"]})
    M.remember_pricing_rule(mem, M.MAX_DISCOUNT_KEY, 15)

    small = X.build_workbook(mem, "commitments", today=TODAY)
    assert small.sheetnames == ["Commitments"]

    full = X.build_workbook(mem, "all", today=TODAY)
    assert full.sheetnames == ["Commitments", "Flagged promises",
                               "Clients", "Policies", "Audit trail"]

    path = tmp_path / "out.xlsx"
    full.save(path)
    assert load_workbook(path)["Commitments"].max_row == 4  # header + 3


def test_kept_commitments_are_not_exported(mem):
    _fixture(mem)
    rows = X.commitment_rows(mem, today=TODAY)
    M.complete_commitment(mem, rows[0]["id"])
    assert len(X.commitment_rows(mem, today=TODAY)) == len(rows) - 1
