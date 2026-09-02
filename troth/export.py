"""Excel export.

Everything here is a read of Sibyl at the moment the button is pressed —
there is no export cache and no snapshot table. Two scopes:

    commitments   what is owed, filtered by due date
    all           every tier, one sheet each

The date filter only applies to the commitments sheet, because it is the
only data with a date that means anything. Sibyl's COLD journal is
timestamped by when a line was *written*, which is not when anything is
due; filtering the audit trail by "upcoming" would be meaningless.
Undated commitments are never filtered out — a commitment with no
deadline is the one most likely to slip, so hiding it because it fails a
date range would defeat the point.
"""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import commitments as C
from . import memory as M

HEADER_FILL = PatternFill("solid", fgColor="0B1724")
HEADER_FONT = Font(color="F3F1EB", bold=True, size=10)
URGENCY_FILL = {
    "OVERDUE": PatternFill("solid", fgColor="FFD9CC"),
    "DUE": PatternFill("solid", fgColor="FFF0CC"),
    "UNDATED": PatternFill("solid", fgColor="EDEDED"),
}


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int]):
    ws = wb.create_sheet(title) if wb.sheetnames != ["Sheet"] else wb.active
    ws.title = title
    ws.append(headers)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def commitment_rows(memory, date_from: date | None = None,
                    date_to: date | None = None, today: date | None = None) -> list[dict]:
    """Open commitments, optionally within a due-date window."""
    today = today or date.today()
    rows = []
    for rec in M.list_commitments(memory, status=M.OPEN):   # SIBYL WARM
        body = rec.get("body") or {}
        due = C.parse_iso(body.get("due"))
        precision = body.get("precision", "")

        if due is not None:
            if date_from and due < date_from:
                continue
            if date_to and due > date_to:
                continue

        rows.append({
            "id": rec.get("name"),
            "client": body.get("deal", ""),
            "commitment": body.get("text", ""),
            "made_by": body.get("made_by") or "unknown",
            "said": body.get("phrase") or "",
            "due": due,
            "precision": precision,
            "urgency": C.status_for(due, precision, today),
            "detail": C.describe(due, precision, today),
        })

    # Overdue first, then soonest. Undated sort last but are never dropped.
    return sorted(rows, key=lambda r: (r["due"] is None, r["due"] or date.max))


def build_workbook(memory, scope: str = "commitments",
                   date_from: date | None = None, date_to: date | None = None,
                   today: date | None = None) -> Workbook:
    today = today or date.today()
    wb = Workbook()

    ws = _sheet(wb, "Commitments",
                ["Due", "Urgency", "Client", "Commitment", "Said", "By", "Detail"],
                [12, 11, 18, 62, 18, 14, 22])
    for r in commitment_rows(memory, date_from, date_to, today):
        ws.append([
            r["due"].isoformat() if r["due"] else "—",
            r["urgency"], r["client"], r["commitment"],
            r["said"], r["made_by"], r["detail"],
        ])
        fill = URGENCY_FILL.get(r["urgency"])
        if fill:
            for col in range(1, 8):
                ws.cell(row=ws.max_row, column=col).fill = fill

    if scope != "all":
        return wb

    ws = _sheet(wb, "Flagged promises",
                ["Status", "Client", "Promise", "By", "Why it was flagged"],
                [11, 18, 62, 14, 46])
    for p in memory.list_entities("promise"):               # SIBYL WARM
        body = p.get("body") or {}
        ws.append([p.get("status", ""), body.get("deal", ""), body.get("text", ""),
                   body.get("made_by") or "", body.get("reason", "")])

    ws = _sheet(wb, "Clients", ["Client", "Status", "Facts"], [22, 12, 96])
    for c in memory.list_entities("client"):                # SIBYL WARM
        body = c.get("body") or {}
        ws.append([c.get("name", ""), c.get("status", ""),
                   " · ".join(body.get("facts", []))])

    ws = _sheet(wb, "Policies", ["Rule", "Value", "Updated"], [34, 14, 22])
    for p in M.list_policies(memory):                       # SIBYL REFERENCE
        ws.append([p["key"], p["value"], (p.get("updated") or "")[:19]])

    ws = _sheet(wb, "Audit trail", ["When", "Kind", "Client", "Entry"],
                [22, 22, 18, 86])
    for e in M.timeline(memory, limit=1000):                # SIBYL COLD
        extra = e.get("extra") or {}
        ws.append([(e.get("ts") or "")[:19].replace("T", " "),
                   extra.get("kind", ""), extra.get("deal", ""),
                   (e.get("acted") or [""])[0]])

    return wb


def to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
